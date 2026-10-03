"""Assemble the SSC-style MERAS flowcharts into docs/MERAS-System-Flowcharts.docx (one chart per page)."""
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

sys.path.insert(0, str(Path(__file__).parent))
import generate_meras_flowcharts as g  # noqa: E402  (registers all charts)

MAX_W, MAX_H = 6.3, 8.2  # inches available on a Letter/A4 page with 1" margins

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)

doc.add_heading("System Flowchart", level=1)
doc.add_paragraph(
    "The flowcharts below trace the MERAS web and mobile system from the moment the URL is entered. "
    "An oval marks the start or end, a parallelogram an input, a rectangle a page or process, a diamond a "
    "decision, a cylinder a database table, a circle an on-page connector, and a downward pentagon an "
    "off-page connector that continues on the chart bearing the same code."
)

from PIL import Image  # noqa: E402

first = True
for slug, title, _fn, _kw in g.CHARTS:
    path = g.OUT / f"fc-{slug}.png"
    if not path.exists():
        raise SystemExit(f"missing {path}; run generate_meras_flowcharts.py first")
    w, h = Image.open(path).size
    ratio = w / h
    if MAX_W / ratio <= MAX_H:
        width = Inches(MAX_W)
    else:
        width = Inches(MAX_H * ratio)
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=width)
    cap = doc.add_paragraph(title)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].bold = True

out = Path(__file__).parent / "MERAS-System-Flowcharts.docx"
doc.save(out)
print("saved", out)
