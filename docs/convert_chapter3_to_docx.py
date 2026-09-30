from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE_PATH = Path(__file__).with_name("chapter-3-methodology.md")
OUTPUT_PATH = Path(__file__).with_name("Meras_Chapter_3_RAD_Methodology_With_Diagrams.docx")
FIGURE_PATHS = {
    "Figure 1": "figure-01-rad-lifecycle.png",
    "Figure 2": "figure-02-gantt-chart.png",
    "Figure 3": "figure-03-context-flow-diagram.png",
    "Figure 4": "figure-04-level-1-dfd.png",
    "Figure 5": "figure-05-customer-support-dfd.png",
    "Figure 6": "figure-06-inventory-dfd.png",
    "Figure 7": "figure-07-pos-dfd.png",
    "Figure 8": "figure-08-system-architecture.png",
    "Figure 9": "figure-09-entity-relationship-diagram.png",
}


def set_cell_shading(cell, fill):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def add_inline_text(paragraph, value):
    chunks = re.split(r"(\*\*.*?\*\*)", value)
    for chunk in chunks:
        if not chunk:
            continue
        bold = chunk.startswith("**") and chunk.endswith("**")
        run = paragraph.add_run(chunk[2:-2] if bold else chunk)
        run.bold = bold
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.size = Pt(12)


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)


def add_heading(document, value, level):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12 if level > 1 else 0)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.5
    if level == 1:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        size = 14
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        size = 12
    run = paragraph.add_run(value)
    run.bold = True
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)


def add_body_paragraph(document, value, is_figure=False):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(0)
    if is_figure:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(6)
        run = paragraph.add_run(value)
        run.italic = True
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.size = Pt(12)
        return
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0.5)
    add_inline_text(paragraph, value)


def add_bullet(document, value):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    add_inline_text(paragraph, value)


def add_figure(document, caption):
    figure_number = caption.split(".", 1)[0]
    figure_path = Path(__file__).with_name("chapter-3-diagrams") / FIGURE_PATHS[figure_number]
    if not figure_path.exists():
        raise FileNotFoundError(f"Missing figure asset: {figure_path}")
    document.add_page_break()
    figure_paragraph = document.add_paragraph()
    figure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_paragraph.paragraph_format.space_after = Pt(4)
    figure_paragraph.add_run().add_picture(str(figure_path), width=Inches(6.1))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(6)
    caption_run = caption_paragraph.add_run(caption)
    caption_run.italic = True
    caption_run.font.name = "Times New Roman"
    caption_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    caption_run.font.size = Pt(11)


def add_table(document, rows):
    cells = [[item.strip() for item in row.strip().strip("|").split("|")] for row in rows]
    table = document.add_table(rows=0, cols=len(cells[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, values in enumerate(cells):
        table_row = table.add_row()
        for column_index, value in enumerate(values):
            cell = table_row.cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(value)
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            run.font.size = Pt(10)
            if row_index == 0:
                set_cell_shading(cell, "D9EAD3")
                run.bold = True
    document.add_paragraph()


def build_document():
    document = Document()
    configure_document(document)
    lines = SOURCE_PATH.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                candidate = lines[index].strip()
                if not re.fullmatch(r"\|\s*[-: ]+(\|\s*[-: ]+)+\|?", candidate):
                    table_lines.append(candidate)
                index += 1
            add_table(document, table_lines)
            continue
        if line.startswith("### "):
            add_heading(document, line[4:], 3)
        elif line.startswith("## "):
            add_heading(document, line[3:], 2)
        elif line.startswith("# "):
            add_heading(document, line[2:], 1)
        elif line.startswith("- "):
            add_bullet(document, line[2:])
        else:
            is_figure = line.startswith("**Figure ") and line.endswith("**")
            if is_figure:
                add_figure(document, line[2:-2])
            else:
                add_body_paragraph(document, line)
        index += 1
    document.save(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
